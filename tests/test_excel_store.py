from openpyxl import load_workbook

from excel_store import COLUMNS, JobStore, split_list
from workspace import Workspace, desktop_dir


def test_new_workbook_has_jobs_sheet_and_all_columns(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.save()
    ws = load_workbook(tmp_path / "jobs.xlsx")["Jobs"]
    assert [c.value for c in ws[1]] == COLUMNS


def test_upsert_insert_then_update_by_job_id(tmp_path):
    path = tmp_path / "jobs.xlsx"
    store = JobStore(path)
    assert store.upsert("a1", {"Job Title": "Data Analyst", "Matched Skills": ["SQL", "Python"]}) == "inserted"
    store.save()
    store = JobStore(path)
    assert store.upsert("a1", {"Match Score": 9}) == "updated"
    store.save()
    row = JobStore(path).get("a1")
    assert row["Job Title"] == "Data Analyst" and row["Match Score"] == 9
    assert split_list(row["Matched Skills"]) == ["SQL", "Python"]
    assert len(JobStore(path)) == 1 and row["Last Updated"]


def test_existing_rows_and_user_edits_preserved(tmp_path):
    path = tmp_path / "jobs.xlsx"
    store = JobStore(path)
    store.upsert("a1", {"Job Title": "A"})
    store.upsert("b2", {"Job Title": "B"})
    store.save()
    # user edits in Excel: application status, a personal notes column, column order unchanged
    wb = load_workbook(path)
    ws = wb["Jobs"]
    col = {c.value: c.column for c in ws[1]}
    ws.cell(row=2, column=col["Application Status"], value="Applied")
    ws.cell(row=1, column=ws.max_column + 1, value="My Notes")
    ws.cell(row=2, column=ws.max_column, value="call back Friday")
    wb.save(path)

    store = JobStore(path)
    store.upsert("a1", {"Processing Status": "ready_for_review", "Application Status": "Not Applied"})
    store.upsert("c3", {"Job Title": "C"})
    store.save()
    store = JobStore(path)
    assert {r["Job ID"] for r in store.all()} == {"a1", "b2", "c3"}
    assert store.get("a1")["Application Status"] == "Applied"          # user-owned, never overwritten
    assert store.get("a1")["My Notes"] == "call back Friday"
    assert store.get("c3")["Application Status"] == "Not Applied"


def test_ref_id_kept_as_exact_text_and_null_allowed(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Source Reference ID": "0004467378631"})
    store.upsert("b2", {"Source Reference ID": None})
    store.save()
    store = JobStore(tmp_path / "jobs.xlsx")
    assert store.get("a1")["Source Reference ID"] == "0004467378631"
    assert store.get("b2")["Source Reference ID"] is None


def test_scraped_formula_text_is_stored_as_text(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Job Description": "=HYPERLINK(\"http://evil\")"})
    store.save()
    cell = load_workbook(tmp_path / "jobs.xlsx")["Jobs"]["J2"]
    assert cell.data_type == "s"


def test_long_description_truncated_to_excel_limit(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Job Description": "x" * 40000})
    store.save()
    assert len(JobStore(tmp_path / "jobs.xlsx").get("a1")["Job Description"]) == 32000


def test_save_keeps_backup(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {})
    store.save()
    store.upsert("b2", {})
    store.save()
    assert (tmp_path / "jobs.xlsx.bak").is_file()


def test_dedup_index_shape(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Source Reference ID": "1", "Job Title": "T", "Company": "C", "Job URL": "u",
                        "Date Discovered": "2026-10-13", "Processing Status": "discovered"})
    idx = store.dedup_index()
    assert idx["a1"]["title"] == "T" and idx["a1"]["company"] == "C" and idx["a1"]["first_seen_at"] == "2026-10-13"


def test_desktop_resolved_from_home_dir(tmp_path):
    (tmp_path / "OneDrive" / "Desktop").mkdir(parents=True)
    assert desktop_dir(tmp_path) == tmp_path / "OneDrive" / "Desktop"
    (tmp_path / "Desktop").mkdir()
    assert desktop_dir(tmp_path) == tmp_path / "Desktop"


def test_workspace_setup_creates_structure_and_never_overwrites(isolated_workspace):
    ws = Workspace(isolated_workspace).setup()
    for d in (ws.resumes_dir, ws.logs_dir, ws.config_dir):
        assert d.is_dir()
    assert (ws.config_dir / "search.json").is_file() and ws.prd.is_file()
    ws.prd.write_text("my PRD", encoding="utf-8")
    ws.setup()
    assert ws.prd.read_text(encoding="utf-8") == "my PRD"


def test_upsert_none_clears_cell(tmp_path):
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Error Details": "boom"})
    store.upsert("a1", {"Error Details": None})
    store.save()
    assert JobStore(tmp_path / "jobs.xlsx").get("a1")["Error Details"] is None


def test_locked_workbook_raises_clear_error(tmp_path, monkeypatch):
    import pytest

    import excel_store
    JobStore(tmp_path / "jobs.xlsx").save()

    def locked(*a, **k):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(excel_store, "load_workbook", locked)
    with pytest.raises(excel_store.WorkbookLockedError, match="close it in Excel"):
        JobStore(tmp_path / "jobs.xlsx")


def test_dashboard_sheet_built_on_save_and_rebuilt_not_duplicated(tmp_path):
    from openpyxl import load_workbook
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Job Title": "AI Engineer", "Company": "Acme", "Match Score": 9.6,
                        "Processing Status": "verified_top_match", "Job URL": "https://www.linkedin.com/jobs/view/1",
                        "Date Discovered": "2026-10-13"})
    store.upsert("b2", {"Match Score": 7.25, "Processing Status": "rejected_low_fit", "Date Discovered": "2026-10-13"})
    store.upsert("c3", {"Processing Status": "excluded_experience",
                        "Error Details": "LinkedIn seniority 'Mid-Senior level' is above fresher level"})
    store.save()
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("d4", {"Processing Status": "excluded_staffing", "Error Details": "staffing/recruiting company name: X HR"})
    store.save()
    wb = load_workbook(tmp_path / "jobs.xlsx")
    assert wb.sheetnames.count("Dashboard") == 1 and "Jobs" in wb.sheetnames
    dash = wb["Dashboard"]
    cells = {c.value for row in dash.iter_rows() for c in row if c.value is not None}
    assert "1 / 200  (0%)" in cells                         # goal progress
    assert "AI Engineer" in cells and "https://www.linkedin.com/jobs/view/1" in cells
    assert "LinkedIn seniority: Mid-Senior level" in cells and "Staffing / recruiting company" in cells
    assert len(dash._charts) == 2
    assert JobStore(tmp_path / "jobs.xlsx").get("a1")["Match Score"] == 9.6   # Jobs data untouched


def test_control_characters_in_scraped_text_are_stripped(tmp_path):
    """Regression: a LinkedIn description with control chars crashed the whole save (IllegalCharacterError)."""
    store = JobStore(tmp_path / "jobs.xlsx")
    store.upsert("a1", {"Job Description": "About\x0bWeRize\x02 hiring\tSQL\nanalysts", "Company": "We\x1fRize"})
    store.save()
    row = JobStore(tmp_path / "jobs.xlsx").get("a1")
    assert row["Job Description"] == "AboutWeRize hiring\tSQL\nanalysts" and row["Company"] == "WeRize"
