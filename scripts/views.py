"""'Master' view sheet in jobs.xlsx: every open opening with Match Score >= 8, from all platforms.

The daily run only fetches jobs posted in the last 24 hours; the ones that score 8+ land here and
stay, so the Master tab grows day by day. Closed or expired openings, and the ones the user marked
'Not Working' (a fake or dead listing), are left out; they stay in the Jobs sheet. It is rebuilt
from the Jobs sheet on every save, and the Jobs sheet (the agent's full reference data) keeps
every row.

The columns you fill in (Applied?, Company Domain, Email Format, Contact Emails) are copied back to
the Jobs sheet on the next run, so they survive the rebuild. Company Domain, Email Format and
Contact Emails belong to the company: typed once, they show on every opening from that company.

Posting times come from the run snapshots in <workspace>/logs (jobs.xlsx has no posted-date column);
without one, Date Discovered is used - searches only return jobs from the last 24 h, so the two are
within a day of each other.
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from excel_store import COMPANY_FIELDS, LEGACY_VIEW_SHEETS, MASTER_SHEET, MASTER_EDITABLE, company_key

PLATFORMS = ("LinkedIn", "Indeed", "Naukri")
MIN_SCORE = 8.0
HIDE_STATUSES = {"closed", "expired"}
IST = timezone(timedelta(hours=5, minutes=30))

LINK, INK, MUTED, HEAD_BG, TOP_BG, EDIT_BG = "2A78D6", "0B0B0B", "8A8984", "E8EEF7", "EAF4EA", "FFF7E0"
THIN = Side(style="thin", color="D9DDE3")

COLUMNS = [("Platform", 11), ("Posted", 17), ("Match Score", 11), ("Status", 26), ("Company", 28),
           ("Employees", 11), ("Job Title", 44), ("Job ID", 18), ("Location", 24), ("Job URL", 40), ("Job Description", 60),
           ("Applied?", 13), ("Company Domain", 22), ("Email Format", 16), ("Emails in JD", 34),
           ("Contact Emails", 40)]
COL = {name: i for i, (name, _) in enumerate(COLUMNS, 1)}
BUILT = [(f"_built:{view}", 4) for view, _ in MASTER_EDITABLE]   # hidden: value at build time, for syncing edits
STATUS_LABEL = {"verified_top_match": "Top match (> 9.0, open)", "top_match_unverified": "Top match (open unconfirmed)",
                "qualified": "Qualified (8 - 9.0)"}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
NOT_EMAIL_TLDS = {"png", "jpg", "jpeg", "gif", "svg", "webp"}       # image@2x.png and similar

_cache: dict[str, tuple[float, dict]] = {}


def source_of(url: str | None) -> str:
    u = (url or "").lower()
    if "indeed." in u:
        return "Indeed"
    if "naukri.com" in u:
        return "Naukri"
    return "LinkedIn"


def emails_in(text: str | None) -> list[str]:
    """Unique email addresses in a job description, lower-cased, in order of appearance."""
    out: list[str] = []
    for m in EMAIL_RE.findall(text or ""):
        m = m.lower().rstrip(".")
        if m.rsplit(".", 1)[-1] not in NOT_EMAIL_TLDS and m not in out:
            out.append(m)
    return out


def _read_snapshot(path: Path) -> dict[str, str]:
    """{job_id: posted timestamp or date} from one run snapshot (cached by mtime)."""
    key, mtime = str(path), path.stat().st_mtime
    if key in _cache and _cache[key][0] == mtime:
        return _cache[key][1]
    from scrape_jobs import make_job_id
    out: dict[str, str] = {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        data = {}
    for hit in data.get("search_hits") or []:                       # LinkedIn search results
        if isinstance(hit, dict) and hit.get("jobId") and hit.get("postedDate"):
            out[make_job_id(str(hit["jobId"]), None)] = str(hit["postedDate"])
    for item in data.get("raw_items") or []:                        # LinkedIn details
        if isinstance(item, dict) and item.get("id") and item.get("datePosted"):
            out.setdefault(make_job_id(str(item["id"]), None), str(item["datePosted"]))
    for job in data.get("jobs") or []:                              # Indeed / Naukri
        if isinstance(job, dict) and job.get("job_id") and (job.get("posted_at") or job.get("posted_date")):
            out[job["job_id"]] = str(job.get("posted_at") or job.get("posted_date"))
    _cache[key] = (mtime, out)
    return out


def posted_dates(logs_dir: Path) -> dict[str, str]:
    dates: dict[str, str] = {}
    if logs_dir.is_dir():
        for snap in sorted(logs_dir.glob("*_discovery.json")) + sorted(logs_dir.glob("*_source_*.json")):
            dates.update(_read_snapshot(snap))
    return dates


def parse_posted(value: str | None) -> datetime | None:
    """A posting time as an aware datetime; a date-only value is midnight India time."""
    if not value:
        return None
    v = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        try:
            return datetime.strptime(v, "%Y-%m-%d").replace(tzinfo=IST)
        except ValueError:
            return None
    try:
        dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
    except ValueError:
        try:
            dt = datetime.strptime(v[:19], "%Y-%m-%d %H:%M:%S").replace(tzinfo=IST)   # Naukri text form
        except ValueError:
            return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def master_rows(rows: list[dict], posted: dict[str, str]) -> list[dict]:
    """Open 8+ openings from every platform, newest posting first, then highest score."""
    out = []
    for r in rows:
        score = r.get("Match Score")
        if not isinstance(score, (int, float)) or isinstance(score, bool) or score < MIN_SCORE:
            continue
        status = str(r.get("Processing Status") or "")
        if status in HIDE_STATUSES or status.startswith("excluded_") or r.get("Application Status") == "Not Working":
            continue
        out.append({**r, "_posted": posted.get(str(r.get("Job ID"))) or r.get("Date Discovered")})
    oldest = datetime.min.replace(tzinfo=timezone.utc)
    out.sort(key=lambda r: (str(r.get("Company") or "").lower()))
    out.sort(key=lambda r: (parse_posted(r["_posted"]) or oldest, r["Match Score"]), reverse=True)
    return out


def company_values(rows: list[dict]) -> dict[str, dict[str, str]]:
    """{company key: {field: value}} for the company-level fields, from any row that has one."""
    out: dict[str, dict[str, str]] = {}
    for r in rows:
        key = company_key(r.get("Company"))
        for field in COMPANY_FIELDS:
            if key and r.get(field) not in (None, ""):
                out.setdefault(key, {}).setdefault(field, r[field])
    return out


def _fmt_posted(value: str | None) -> str | None:
    if not value:
        return None
    v = str(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        return v
    try:
        dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
        return dt.astimezone(IST).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return v[:16]


def build_views(wb, rows: list[dict], logs_dir: Path | None, now: datetime | None = None) -> dict[str, int]:
    now = now or datetime.now(timezone.utc)
    posted = posted_dates(logs_dir) if logs_dir else {}
    for name in (MASTER_SHEET, *LEGACY_VIEW_SHEETS):     # per-platform tabs were replaced by Master
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet(MASTER_SHEET, 1)                 # right after the Dashboard
    ws.sheet_view.showGridLines = False
    items = master_rows(rows, posted)
    by_company = company_values(rows)

    ws["A1"] = f"Master: open openings scoring 8+ from LinkedIn, Indeed and Naukri ({len(items)})"
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws["A2"] = (f"Built {now.astimezone(IST).strftime('%d %b %Y, %H:%M')} IST from the Jobs sheet (full data kept "
                f"there). Rebuilt on every save. Fill the shaded columns (Applied?, Company Domain, Email Format, "
                f"Contact Emails); they are copied to the Jobs sheet on the next run, and the company columns "
                f"apply to every opening from that company.")
    ws["A2"].font = Font(size=9, color=MUTED)
    editable = {view for view, _ in MASTER_EDITABLE}
    for c, (name, width) in enumerate(COLUMNS + BUILT, 1):
        cell = ws.cell(row=4, column=c, value=name)
        cell.font = Font(bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=EDIT_BG if name in editable else HEAD_BG)
        cell.border = Border(bottom=THIN)
        ws.column_dimensions[cell.column_letter].width = width
        if name.startswith("_built:"):
            ws.column_dimensions[cell.column_letter].hidden = True

    for r_i, r in enumerate(items, 5):
        status = r.get("Processing Status")
        company = by_company.get(company_key(r.get("Company")), {})
        mine = {view: (r.get(field) if r.get(field) not in (None, "") or field not in COMPANY_FIELDS
                       else company.get(field)) for view, field in MASTER_EDITABLE}
        found = emails_in(r.get("Job Description"))
        vals = [source_of(r.get("Job URL")), _fmt_posted(r["_posted"]), r.get("Match Score"),
                STATUS_LABEL.get(status, status), r.get("Company"), r.get("Employees"), r.get("Job Title"),
                r.get("Job ID"),
                r.get("Location"), r.get("Job URL"), r.get("Job Description"), mine["Applied?"],
                mine["Company Domain"], mine["Email Format"], "; ".join(found) or None, mine["Contact Emails"],
                *(mine[view] for view, _ in MASTER_EDITABLE)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r_i, column=c, value=v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"                      # scraped text must never become a formula
            cell.font = Font(color=INK)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if status == "verified_top_match":
                cell.fill = PatternFill("solid", fgColor=TOP_BG)
        ws.cell(row=r_i, column=3).number_format = "0.00"
        if r.get("Job URL"):
            link = ws.cell(row=r_i, column=COL["Job URL"])
            link.hyperlink = r["Job URL"]
            link.font = Font(color=LINK, underline="single")
    if not items:
        ws.cell(row=5, column=1, value="No open openings scoring 8+ yet.").font = Font(italic=True, color=MUTED)
    ws.freeze_panes = "G5"                                # platform .. employees stay visible while scrolling
    last = len(COLUMNS)
    ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=last).column_letter}{max(5, 4 + len(items))}"
    if items:
        from excel_store import add_choice_dropdown
        applied = ws.cell(row=4, column=COL["Applied?"]).column_letter
        add_choice_dropdown(ws, f"{applied}5:{applied}{4 + len(items)}")
    return {MASTER_SHEET: len(items)}
