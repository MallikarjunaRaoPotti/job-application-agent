"""Indeed / Naukri normalization, multi-source pipeline, and the Master view tab."""
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

import run_pipeline as rp
from excel_store import JobStore
from scrape_jobs import make_job_id
from sources import normalize_indeed, normalize_naukri
from test_pipeline import LONG_DESC, RAW, FakeLLM, evaluation, rows, top_eval
from views import emails_in, parse_posted, source_of
from workspace import get_workspace

NOW = datetime.now(timezone.utc)
IST = timezone(timedelta(hours=5, minutes=30))


def indeed_item(key="a1b2c3d4e5f60718", expired=False, published=None, title="Junior Data Analyst"):
    return {"key": key, "title": title, "url": f"https://in.indeed.com/viewjob?jk={key}",
            "employer": {"name": "Acme Analytics", "companyPageUrl": "https://in.indeed.com/cmp/Acme"},
            "location": {"city": "Pune", "countryName": "India"}, "description": {"text": LONG_DESC},
            "jobTypes": {"x": "Full-time"},
            "baseSalary": {"min": 300000, "max": 500000, "unitOfWork": "YEAR", "currencyCode": "INR"},
            "datePublished": published or (NOW - timedelta(hours=3)).isoformat(), "expired": expired,
            "attributes": {"a": "SQL", "b": "Power BI"}}


def naukri_item(ref="131026048999", consultant=False, created=None, summary=None):
    created = created if created is not None else (NOW - timedelta(hours=2)).timestamp() * 1000
    return {"id": ref, "title": "Data Analyst - Fresher",
            "url": f"https://www.naukri.com/job-listings-data-analyst-{ref}?src=x",
            "company": {"name": "Beta Corp", "jobsUrl": "https://www.naukri.com/beta-jobs"},
            "description": {"full": summary or ("<p>Analyse sales data with SQL and Power BI, build dashboards, "
                                                "clean and validate data, and report weekly insights to the business team.</p>")},
            "experience": {"text": "0-1 Yrs"}, "salary": {"label": "3-4 Lacs"}, "locations": [{"label": "Hyderabad"}],
            "skills": {"preferred": ["SQL"], "other": ["Power BI"]}, "education": {"ug": ["B.Tech/B.E."]},
            "role": "Data Analyst", "industry": "IT Services", "consultant": consultant, "createdDate": created,
            "employmentType": "Full Time, Permanent", "wfhType": "0"}


# ---- normalization

def test_indeed_normalized_with_namespaced_id():
    j = normalize_indeed(indeed_item(), NOW)
    assert j["ref_id"] == "a1b2c3d4e5f60718" and j["source"] == "Indeed"
    assert j["job_id"] == make_job_id("a1b2c3d4e5f60718", None, "indeed") != make_job_id("a1b2c3d4e5f60718", None)
    assert j["link"] == "https://in.indeed.com/viewjob?jk=a1b2c3d4e5f60718" and j["location"] == "Pune, India"
    assert normalize_indeed(indeed_item(expired=True), NOW)["job_status"] == "expired"
    assert normalize_indeed({"key": "a1b2c3d4", "title": "x", "employer": {}}, NOW) is None


def test_naukri_normalized_with_structured_facts_and_staffing_flag():
    j = normalize_naukri(naukri_item(), NOW)
    assert j["source"] == "Naukri" and j["link"].endswith("-131026048999") and "?" not in j["link"]
    assert j["experience_required"] == "0-1 Yrs" and "Key skills: SQL" in j["description"]
    assert j["description_core_chars"] < 300
    assert normalize_naukri(naukri_item(consultant=True), NOW)["company_industry"] == "Staffing and Recruiting"


def test_source_of_url():
    assert source_of("https://in.indeed.com/viewjob?jk=1") == "Indeed"
    assert source_of("https://www.naukri.com/job-listings-x-1") == "Naukri"
    assert source_of("https://www.linkedin.com/jobs/view/1") == "LinkedIn"


# ---- posting times + emails

def test_parse_posted():
    assert parse_posted("2026-10-14T05:00:00Z") == datetime(2026, 10, 14, 5, tzinfo=timezone.utc)
    assert parse_posted("2026-10-14") == datetime(2026, 10, 14, tzinfo=IST)
    assert parse_posted("2026-10-14 09:30:00") == datetime(2026, 10, 14, 9, 30, tzinfo=timezone.utc)
    assert parse_posted(None) is None and parse_posted("yesterday") is None


def test_emails_in_description():
    text = "Send CV to HR@Acme.com or hr@acme.com. Logo: banner@2x.png; also jobs.india@acme.co.in."
    assert emails_in(text) == ["hr@acme.com", "jobs.india@acme.co.in"]
    assert emails_in(None) == []


# ---- pipeline + views

def run_multi(llm, source_items, **kw):
    ws = get_workspace()
    cfg = {"sources": {"indeed": {"enabled": True}, "naukri": {"enabled": True}}}
    stats = rp.run(ws, run_id=kw.pop("run_id", "r1"), client_factory=llm.factory, search_config=cfg,
                   raw_items=kw.pop("raw_items", []), source_items=source_items, tailor=False, **kw)
    return ws, stats


def _master(wb):
    """Master rows as {header: value} dicts."""
    s = wb["Master"]
    head = [c.value for c in s[4]]
    return [dict(zip(head, (c.value for c in s[r]))) for r in range(5, s.max_row + 1) if s.cell(row=r, column=7).value]


def test_indeed_and_naukri_jobs_flow_through_screen_score_and_views():
    items = {"indeed": [indeed_item(), indeed_item(key="ffff0000eeee1111", title="Data Analyst Intern")],
             "naukri": [naukri_item(), naukri_item(ref="131026048111", consultant=True)]}
    llm = FakeLLM(evaluations=[top_eval(), evaluation((3, 3, 3, 3)), evaluation()])
    checked = []
    ws, stats = run_multi(llm, items, open_checker=lambda ids: checked.extend(ids) or {})
    assert stats["sources"]["indeed"]["eligible"] == 2 and stats["sources"]["naukri"]["eligible"] == 2
    assert stats["excluded_staffing"] == 1 and llm.calls["submit_evaluation"] == 3
    assert checked == []                                  # LinkedIn open-check never sees Indeed/Naukri IDs
    data = rows(ws)
    assert sum(1 for r in data.values() if r["Processing Status"] == "verified_top_match") == 1

    wb = load_workbook(ws.jobs_xlsx)
    assert wb.sheetnames == ["Dashboard", "Master", "Jobs"]
    shown = _master(wb)
    assert len(shown) == 2                                             # only the two 8+ jobs
    assert {m["Platform"] for m in shown} <= {"Indeed", "Naukri"}
    assert len(data) == 4                                              # Jobs sheet keeps everything


def test_master_keeps_older_8plus_postings_newest_first():
    items = {"indeed": [indeed_item(published=(NOW - timedelta(hours=40)).isoformat()),
                        indeed_item(key="ffff0000eeee1111", title="Data Analyst Intern")], "naukri": []}
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval(), top_eval()]), items)
    titles = [m["Job Title"] for m in _master(load_workbook(ws.jobs_xlsx))]
    assert titles == ["Data Analyst Intern", "Junior Data Analyst"]    # 40 h old stays, listed after the newer one


def test_master_hides_closed_and_low_scores():
    items = {"indeed": [indeed_item(), indeed_item(key="ffff0000eeee1111", title="Data Analyst Intern")], "naukri": []}
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval(), evaluation((3, 3, 3, 3))]), items)
    store = JobStore(ws.jobs_xlsx)
    jid = next(r["Job ID"] for r in store.all() if r["Job Title"] == "Junior Data Analyst")
    assert [m["Job ID"] for m in _master(load_workbook(ws.jobs_xlsx))] == [jid]
    store.upsert(jid, {"Processing Status": "closed"})
    store.save()
    assert _master(load_workbook(ws.jobs_xlsx)) == [] and len(rows(ws)) == 2


def test_rerun_dedups_indeed_and_naukri():
    items = {"indeed": [indeed_item()], "naukri": [naukri_item()]}
    ws, _ = run_multi(FakeLLM(evaluations=[evaluation(), evaluation()]), items)
    _, stats = run_multi(FakeLLM(), items, run_id="r2")
    assert stats["sources"]["indeed"]["previously_seen"] == 1 and stats["sources"]["naukri"]["previously_seen"] == 1
    assert len(JobStore(ws.jobs_xlsx)) == 2


def test_linkedin_still_works_alongside_sources():
    ws, stats = run_multi(FakeLLM(evaluations=[evaluation(), evaluation((3, 3, 3, 3))]),
                          {"indeed": [], "naukri": []}, raw_items=RAW)
    assert stats["discovery"]["eligible"] == 2 and stats["evaluated"] == 2


# ---- Application Status dropdowns

def _dv_ranges(ws):
    return [(str(dv.sqref), dv.formula1) for dv in ws.data_validations.dataValidation]


def test_application_status_dropdowns_on_jobs_and_views():
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval()]), {"indeed": [indeed_item()], "naukri": []})
    wb = load_workbook(ws.jobs_xlsx)
    choices = '"Not Applied,Applied,Not Working"'
    assert any(f == choices and r.startswith("P2:") for r, f in _dv_ranges(wb["Jobs"]))      # Application Status
    assert any(r in ("K5", "K5:K5") and f == choices for r, f in _dv_ranges(wb["Master"]))
    assert len(_dv_ranges(wb["Jobs"])) == 1                                                   # not duplicated on resave


def test_choice_made_on_view_is_copied_to_jobs():
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval()]), {"indeed": [indeed_item()], "naukri": []})
    wb = load_workbook(ws.jobs_xlsx)
    wb["Master"]["K5"] = "Applied"
    wb.save(ws.jobs_xlsx)
    store = JobStore(ws.jobs_xlsx)
    assert store.view_edits_synced == 1
    store.save()
    jid = load_workbook(ws.jobs_xlsx)["Master"]["G5"].value
    assert JobStore(ws.jobs_xlsx).get(jid)["Application Status"] == "Applied"
    assert load_workbook(ws.jobs_xlsx)["Master"]["K5"].value == "Applied"                      # survives the rebuild


def test_not_working_choice_removes_job_from_master_on_next_run():
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval()]), {"indeed": [indeed_item()], "naukri": []})
    wb = load_workbook(ws.jobs_xlsx)
    wb["Master"]["K5"] = "Not Working"
    wb.save(ws.jobs_xlsx)
    JobStore(ws.jobs_xlsx).save()                                   # the next run: sync, then rebuild
    assert _master(load_workbook(ws.jobs_xlsx)) == []
    assert next(iter(rows(ws).values()))["Application Status"] == "Not Working"   # kept in Jobs


def test_stale_view_never_overwrites_newer_jobs_edit():
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval()]), {"indeed": [indeed_item()], "naukri": []})
    wb = load_workbook(ws.jobs_xlsx)
    jobs = wb["Jobs"]
    col = [c.value for c in jobs[1]].index("Application Status") + 1
    jobs.cell(row=2, column=col, value="Not Working")              # edited on Jobs; the view still says Not Applied
    wb.save(ws.jobs_xlsx)
    store = JobStore(ws.jobs_xlsx)
    assert store.view_edits_synced == 0
    store.save()
    assert next(iter(JobStore(ws.jobs_xlsx).all()))["Application Status"] == "Not Working"


# ---- company columns on Master

def test_company_domain_typed_once_applies_to_every_job_of_that_company():
    items = {"indeed": [indeed_item(), indeed_item(key="ffff0000eeee1111", title="Data Analyst Intern")],
             "naukri": [naukri_item()]}
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval(), top_eval(), top_eval()]), items)
    wb = load_workbook(ws.jobs_xlsx)
    head = [c.value for c in wb["Master"][4]]
    row = next(r for r in range(5, 8) if wb["Master"].cell(row=r, column=5).value == "Acme Analytics")
    wb["Master"].cell(row=row, column=head.index("Company Domain") + 1, value=" https://www.Acme.com/careers ")
    wb["Master"].cell(row=row, column=head.index("Email Format") + 1, value="first.last")
    wb.save(ws.jobs_xlsx)
    store = JobStore(ws.jobs_xlsx)
    assert store.view_edits_synced == 2
    store.save()
    data = rows(ws).values()
    assert [r["Company Domain"] for r in data if r["Company"] == "Acme Analytics"] == ["acme.com", "acme.com"]
    assert [r["Company Domain"] for r in data if r["Company"] == "Beta Corp"] == [None]
    shown = _master(load_workbook(ws.jobs_xlsx))
    assert {(m["Company"], m["Company Domain"], m["Email Format"]) for m in shown} == {
        ("Acme Analytics", "acme.com", "first.last"), ("Beta Corp", None, None)}


def test_later_job_from_known_company_shows_its_domain():
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval()]), {"indeed": [indeed_item()], "naukri": []})
    store = JobStore(ws.jobs_xlsx)
    jid = next(iter(store.all()))["Job ID"]
    store.ws.cell(row=2, column=store.col["Company Domain"], value="acme.com")    # typed on the Jobs sheet
    store.save()
    run_multi(FakeLLM(evaluations=[top_eval()]),
              {"indeed": [indeed_item(key="ffff0000eeee1111", title="Data Analyst Intern")], "naukri": []}, run_id="r2")
    shown = _master(load_workbook(ws.jobs_xlsx))
    assert len(shown) == 2 and all(m["Company Domain"] == "acme.com" for m in shown)
    assert JobStore(ws.jobs_xlsx).get(jid)["Company Domain"] == "acme.com"


def test_choice_left_on_old_platform_tab_is_kept_when_master_replaces_it():
    ws, _ = run_multi(FakeLLM(evaluations=[top_eval()]), {"indeed": [indeed_item()], "naukri": []})
    jid = next(iter(rows(ws)))
    wb = load_workbook(ws.jobs_xlsx)
    old = wb.create_sheet("Indeed")                     # the layout the previous version wrote
    for c, name in enumerate(["Applied?", "Job ID", "_built"], 1):
        old.cell(row=4, column=c, value=name)
    old.append(["Applied", jid, "Not Applied"])
    wb.save(ws.jobs_xlsx)
    store = JobStore(ws.jobs_xlsx)
    assert store.view_edits_synced == 1
    store.save()
    assert "Indeed" not in load_workbook(ws.jobs_xlsx).sheetnames
    assert JobStore(ws.jobs_xlsx).get(jid)["Application Status"] == "Applied"
