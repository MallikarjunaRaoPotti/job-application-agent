"""'Dashboard' sheet in jobs.xlsx: a read-only summary rebuilt from the Jobs sheet on every save.

Sections: goal progress + KPI tiles, pipeline status (table + bar chart), match-score distribution
(table + column chart), verified top matches (clickable links), why jobs were excluded, and jobs
discovered per day. The sheet is owned by code - it is deleted and redrawn each time, so edits to
it don't survive; the Jobs sheet is never touched here.
"""

from __future__ import annotations

import re
from collections import Counter
from datetime import datetime

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET = "Dashboard"
GOAL = 200
TOP_BAR = 9.0

SERIES = "2A78D6"        # single-series bars (reference palette, categorical slot 1)
INK, INK_2, MUTED = "0B0B0B", "52514E", "8A8984"
TILE_BG, HEAD_BG, RULE = "F3F6FB", "E8EEF7", "D9DDE3"

STATUS_ORDER = [
    ("verified_top_match", "Verified top match (> 9.0, open)"),
    ("top_match_unverified", "Top match, open status unconfirmed"),
    ("qualified", "Qualified (8.0 - 9.0)"),
    ("rejected_low_fit", "Rejected: score below 8"),
    ("excluded_experience", "Excluded: experience / seniority"),
    ("excluded_staffing", "Excluded: staffing agency"),
    ("excluded_red_flag", "Excluded: red flag"),
    ("excluded_small_company", "Excluded: under 50 employees"),
    ("expired", "Expired before scoring"),
    ("closed", "Closed after matching"),
    ("unscored", "Waiting to be scored"),
    ("discovered", "Discovered, not yet scored"),
    ("scoring_error", "Scoring error (will retry)"),
    ("ready_for_review", "Resume ready (earlier phase)"),
]
SCORE_BUCKETS = [("Below 5", 0, 5), ("5 - 6", 5, 6), ("6 - 7", 6, 7), ("7 - 8", 7, 8),
                 ("8 - 9", 8, 9.0000001), ("Above 9.0", 9.0000001, 11)]

THIN = Side(style="thin", color=RULE)


def _head(ws, row, col, labels, widths=None):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = Font(bold=True, color=INK)
        c.fill = PatternFill("solid", fgColor=HEAD_BG)
        c.border = Border(bottom=THIN)
        c.alignment = Alignment(vertical="center")


def _section(ws, row, col, title, note=None):
    ws.cell(row=row, column=col, value=title).font = Font(bold=True, size=13, color=INK)
    if note:
        ws.cell(row=row + 1, column=col, value=note).font = Font(italic=True, size=9, color=MUTED)


def _bar_chart(ws, title, cats, vals, anchor, horizontal, width, height, y_title):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = title
    ch.legend = None                      # single series: the title names it
    ch.style = 2
    ch.height, ch.width = height, width
    ch.gapWidth = 60
    ch.y_axis.majorGridlines = None     # values are in the adjacent table; no redundant axis title
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"   # keep table order top-to-bottom
        ch.y_axis.crosses = "max"                  # ...but keep the value axis at the bottom, clear of the title
    ch.add_data(vals, titles_from_data=False)
    ch.set_categories(cats)
    s = ch.series[0]
    s.graphicalProperties.solidFill = SERIES
    s.graphicalProperties.line.noFill = True
    ws.add_chart(ch, anchor)


def _reason_group(status: str, detail: str) -> str:
    d = (detail or "").lower()
    if status == "excluded_experience":
        if "seniority" in d:
            m = re.search(r"seniority '([^']+)'", detail or "")
            return f"LinkedIn seniority: {m.group(1)}" if m else "LinkedIn seniority above fresher"
        return "Requires more than 1 year"
    if status == "excluded_staffing":
        return "Agency posting for a client" if "behalf" in d else "Staffing / recruiting company"
    if status == "excluded_red_flag":
        for key, label in (("unpaid", "Unpaid work"), ("fee", "Asks for a fee or deposit"),
                           ("too short", "Missing / very short description"), ("company page", "No company page"),
                           ("whatsapp", "Apply-via-WhatsApp")):
            if key in d:
                return label
        return "Other red flag"
    if status == "excluded_small_company":
        return "Company under 50 employees"
    return status


def build_dashboard(wb, rows: list[dict]) -> None:
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)        # first tab
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMN", (52, 12, 4, 36, 30, 14, 12, 13, 46, 4, 30, 10, 4, 4)):
        ws.column_dimensions[col].width = w

    statuses = Counter(r.get("Processing Status") or "unknown" for r in rows)
    scores = [r["Match Score"] for r in rows
              if isinstance(r.get("Match Score"), (int, float)) and not isinstance(r.get("Match Score"), bool)]
    verified = sorted((r for r in rows if r.get("Processing Status") == "verified_top_match"),
                      key=lambda r: -(r.get("Match Score") or 0))
    excluded = sum(v for k, v in statuses.items() if k.startswith("excluded_"))

    # ---------------------------------------------------------- header + KPIs
    ws["A1"] = "Job search dashboard"
    ws["A1"].font = Font(bold=True, size=18, color=INK)
    ws["A2"] = (f"Refreshed {datetime.now().strftime('%d %b %Y, %H:%M')}  ·  built from the Jobs sheet  ·  "
                f"top match = Match Score > {TOP_BAR:g} and still open on LinkedIn")
    ws["A2"].font = Font(size=9, color=MUTED)

    pct = len(verified) / GOAL
    ws["A4"] = f"Goal: {GOAL} verified top matches"
    ws["A4"].font = Font(bold=True, color=INK_2)
    ws["A5"] = f"{len(verified)} / {GOAL}  ({pct:.0%})"
    ws["A5"].font = Font(bold=True, size=16, color=INK)
    filled = min(20, round(pct * 20))
    ws["D5"] = "█" * filled + "░" * (20 - filled)
    ws["D5"].font = Font(size=14, color=SERIES)

    tiles = [("Jobs tracked", len(rows)), ("Scored", len(scores)),
             (f"Verified top (> {TOP_BAR:g})", len(verified)), ("Qualified (8 - 9)", statuses.get("qualified", 0)),
             ("Excluded", excluded), ("Waiting to score", statuses.get("unscored", 0) + statuses.get("discovered", 0))]
    for i, (label, value) in enumerate(tiles):
        r, c = 7 + (i // 3) * 3, (1, 4, 6)[i % 3]
        lab = ws.cell(row=r, column=c, value=label)
        val = ws.cell(row=r + 1, column=c, value=value)
        lab.font, val.font = Font(size=9, color=INK_2), Font(bold=True, size=16, color=INK)
        for cell in (lab, val):
            cell.fill = PatternFill("solid", fgColor=TILE_BG)

    # ---------------------------------------------------------- pipeline status
    top = 14
    _section(ws, top, 1, "Pipeline status", "Where every tracked job currently stands")
    _head(ws, top + 2, 1, ["Status", "Jobs"])
    r = top + 3
    known = {k for k, _ in STATUS_ORDER}
    status_rows = [(label, statuses[k]) for k, label in STATUS_ORDER if statuses.get(k)]
    status_rows += [(k, v) for k, v in statuses.items() if k not in known]
    for label, n in status_rows:
        ws.cell(row=r, column=1, value=label).font = Font(color=INK)
        ws.cell(row=r, column=2, value=n).font = Font(color=INK)
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=r, column=2, value=len(rows)).font = Font(bold=True)
    if status_rows:
        _bar_chart(ws, "Jobs by status",
                   Reference(ws, min_col=1, min_row=top + 3, max_row=top + 2 + len(status_rows)),
                   Reference(ws, min_col=2, min_row=top + 3, max_row=top + 2 + len(status_rows)),
                   f"D{top}", horizontal=True, width=18, height=7.5, y_title="Jobs")

    # ---------------------------------------------------------- score distribution
    top = max(r + 3, top + 18)
    _section(ws, top, 1, "Match score distribution", f"Scored jobs only ({len(scores)}); > {TOP_BAR:g} is the target")
    _head(ws, top + 2, 1, ["Match Score", "Jobs"])
    for i, (label, lo, hi) in enumerate(SCORE_BUCKETS):
        ws.cell(row=top + 3 + i, column=1, value=label).font = Font(color=INK)
        ws.cell(row=top + 3 + i, column=2, value=sum(1 for s in scores if lo <= s < hi)).font = Font(color=INK)
    _bar_chart(ws, "Scored jobs by Match Score",
               Reference(ws, min_col=1, min_row=top + 3, max_row=top + 2 + len(SCORE_BUCKETS)),
               Reference(ws, min_col=2, min_row=top + 3, max_row=top + 2 + len(SCORE_BUCKETS)),
               f"D{top}", horizontal=False, width=18, height=7.5, y_title="Jobs")

    # ---------------------------------------------------------- verified top matches
    top = top + 18
    _section(ws, top, 1, f"Verified top matches ({len(verified)})",
             "Match Score > 9.0 and confirmed open on LinkedIn at the last check - review before applying")
    cols = ["Job Title", "Match Score", "", "Company", "Location", "Work Mode", "Job Type", "Applied?", "Job URL"]
    _head(ws, top + 2, 1, cols)
    r = top + 3
    for row in verified:
        vals = [row.get("Job Title"), row.get("Match Score"), None, row.get("Company"), row.get("Location"),
                row.get("Work Mode"), row.get("Job Type"), row.get("Application Status"), row.get("Job URL")]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + i, value=v)
            c.font = Font(color=INK)
            c.border = Border(bottom=THIN)
        url = row.get("Job URL")
        if url:
            link = ws.cell(row=r, column=9)
            link.hyperlink = url
            link.font = Font(color=SERIES, underline="single")
        ws.cell(row=r, column=2).number_format = "0.00"
        r += 1
    if not verified:
        ws.cell(row=r, column=1, value="None yet").font = Font(italic=True, color=MUTED)
        r += 1

    # ---------------------------------------------------------- by source
    from views import PLATFORMS, source_of
    top = r + 2
    _section(ws, top, 1, "By source", "Every open opening scoring 8+ is on the Master tab (filter by Platform)")
    _head(ws, top + 2, 1, ["Source", "Tracked", "", "Scored", "Score 8+", f"Verified top (> {TOP_BAR:g})"])
    for i, src in enumerate(PLATFORMS):
        mine = [x for x in rows if source_of(x.get("Job URL")) == src]
        sc = [x["Match Score"] for x in mine if isinstance(x.get("Match Score"), (int, float))
              and not isinstance(x.get("Match Score"), bool)]
        vals = [src, len(mine), None, len(sc), sum(1 for v in sc if v >= 8),
                sum(1 for x in mine if x.get("Processing Status") == "verified_top_match")]
        for c, v in enumerate(vals, 1):
            ws.cell(row=top + 3 + i, column=c, value=v).font = Font(color=INK)
    r = top + 3 + len(PLATFORMS)

    # ---------------------------------------------------------- exclusions + discovery by day
    top = r + 2
    _section(ws, top, 1, "Why jobs were excluded", "Screened out before scoring - no LLM cost")
    _head(ws, top + 2, 1, ["Reason", "Jobs"])
    reasons = Counter(_reason_group(row["Processing Status"], row.get("Error Details") or "")
                      for row in rows if str(row.get("Processing Status") or "").startswith("excluded_"))
    r = top + 3
    for reason, n in reasons.most_common():
        ws.cell(row=r, column=1, value=reason).font = Font(color=INK)
        ws.cell(row=r, column=2, value=n).font = Font(color=INK)
        r += 1
    if not reasons:
        ws.cell(row=r, column=1, value="None").font = Font(italic=True, color=MUTED)

    _section(ws, top, 4, "Jobs discovered per day")
    _head(ws, top + 2, 4, ["Date discovered", "Jobs", f"Top (> {TOP_BAR:g})"])
    by_day, top_by_day = Counter(), Counter()
    for row in rows:
        day = str(row.get("Date Discovered") or "unknown")[:10]
        by_day[day] += 1
        if row.get("Processing Status") == "verified_top_match":
            top_by_day[day] += 1
    for i, day in enumerate(sorted(by_day, reverse=True)):
        ws.cell(row=top + 3 + i, column=4, value=day).font = Font(color=INK)
        ws.cell(row=top + 3 + i, column=5, value=by_day[day]).font = Font(color=INK)
        ws.cell(row=top + 3 + i, column=6, value=top_by_day[day]).font = Font(color=INK)
