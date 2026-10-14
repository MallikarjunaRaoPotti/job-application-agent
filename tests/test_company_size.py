"""Companies under 50 employees are excluded before scoring (LinkedIn size, cache, or a trusted Apollo match)."""
from openpyxl import load_workbook

from company_size import (ApolloUnavailable, CompanySizer, employees_label, is_small, linkedin_size, linkedin_slug,
                          name_key)
from excel_store import JobStore
from test_pipeline import FakeLLM, evaluation, rows, run, top_eval
from test_sources import _master

def org(employees, domain, name, linkedin="acme"):   # RAW jobs link linkedin.com/company/acme
    return {"employees": employees, "domain": domain, "name": name, "linkedin": linkedin}


class FakeApollo:
    """name or domain -> Apollo answer; missing means 'no record'; 'DOWN' raises like an HTTP 429."""

    def __init__(self, answers: dict):
        self.answers, self.calls = answers, []

    def __call__(self, name=None, domain=None):
        self.calls.append(domain or name)
        ans = self.answers.get(domain or name)
        if ans == "DOWN":
            raise ApolloUnavailable("Apollo HTTP 429")
        return ans


def sizer(tmp_path, answers, **kw):
    fake = FakeApollo(answers)
    return CompanySizer(tmp_path / "company_cache.json", fake, **kw), fake


def opened(ids):
    return {i: (True, "open") for i in ids}


# ---- parsing + rules

def test_parsing_helpers_and_rule():
    assert linkedin_size("11-50 employees") == (11, 50) and linkedin_size("10,001+ employees") == (10001, None)
    assert linkedin_size("1,001-5,000 employees") == (1001, 5000) and linkedin_size("Self-employed") is None
    assert linkedin_slug("https://in.linkedin.com/company/FlexIntl?trk=x") == "flexintl"
    assert linkedin_slug("https://in.indeed.com/cmp/Openjobs") is None
    assert name_key("Zenwork, Inc.") == "zenwork" == name_key("zenwork") and name_key("EY") == "ey"
    assert is_small(2, None) and is_small(49, None) and not is_small(50, None)
    assert is_small(None, (11, 50)) and not is_small(None, (51, 200)) and not is_small(None, (10001, None))
    assert not is_small(None, None)                                   # unknown is kept


def test_employees_label():
    assert employees_label({"employees": 2}) == "2"
    assert employees_label({"employees": None, "range": [51, 200]}) == "51-200"
    assert employees_label({"employees": None, "range": [10001, None]}) == "10,001+"
    assert employees_label({"employees": None, "range": None}) == "unknown" == employees_label(None)


# ---- trust: a namesake must never exclude a real job

def test_namesake_from_name_lookup_is_not_trusted(tmp_path):
    s, _ = sizer(tmp_path, {"Flex": org(400, "getflex.com", "Flex", linkedin="getflex")})
    res = s.assess("Flex", company_url="https://www.linkedin.com/company/flexintl")
    assert res == {"employees": "unknown", "small": False, "reason": None, "domain": None}


def test_name_lookup_trusted_when_linkedin_page_matches(tmp_path):
    s, _ = sizer(tmp_path, {"Tiny Courses": org(3, "tiny.io", "Tiny Courses", linkedin="tiny-courses")})
    res = s.assess("Tiny Courses", company_url="https://in.linkedin.com/company/tiny-courses")
    assert res["small"] and res["employees"] == "3" and res["domain"] == "tiny.io"
    assert "Apollo: 3 employees" in res["reason"]


def test_indeed_job_trusted_only_on_exact_name(tmp_path):
    s, _ = sizer(tmp_path, {"Openjobs": org(5, "openjobshq.com", "OpenJobs", linkedin="openjobs"),
                      "Nova": org(5, "nova.io", "Nova Labs Global", linkedin="nova")})
    assert s.assess("Openjobs", company_url="https://in.indeed.com/cmp/Openjobs")["small"]
    assert not s.assess("Nova", company_url="https://in.indeed.com/cmp/Nova")["small"]   # different name: unknown


def test_domain_lookup_is_exact_and_overrides_a_name_answer(tmp_path):
    s, fake = sizer(tmp_path, {"Flex": org(400, "getflex.com", "Flex", linkedin="getflex"),
                         "flex.com": org(150000, "flex.com", "Flex", linkedin="flexintl")})
    assert s.assess("Flex")["employees"] == "400"                      # name-only, exact name -> trusted
    res = s.assess("Flex", domain="https://www.flex.com/")             # the user typed the real domain
    assert res["employees"] == "150000" and not res["small"] and fake.calls == ["Flex", "flex.com"]
    s.assess("Flex", domain="flex.com")
    assert fake.calls == ["Flex", "flex.com"]                          # same domain: from the cache


def test_linkedin_size_needs_no_apollo_call(tmp_path):
    s, fake = sizer(tmp_path, {})
    assert "LinkedIn: 11-50 employees" in s.assess("Small Co", "11-50 employees")["reason"]
    assert not s.assess("Big Co", "10,001+ employees")["small"] and fake.calls == []


# ---- cache, outages, cap

def test_each_company_looked_up_once_and_cached_across_runs(tmp_path):
    s, fake = sizer(tmp_path, {"Tiny": org(3, "tiny.io", "Tiny")})
    assert s.assess("Tiny")["small"] and s.assess("tiny ")["small"] and fake.calls == ["Tiny"]
    s.save()
    s2, fake2 = sizer(tmp_path, {})
    assert s2.assess("Tiny")["small"] and fake2.calls == []           # from the cache file


def test_unknown_company_is_kept_and_not_asked_again(tmp_path):
    s, fake = sizer(tmp_path, {})
    assert s.assess("Ghost Ltd")["employees"] == "unknown" and s.assess("Ghost Ltd")["employees"] == "unknown"
    assert fake.calls == ["Ghost Ltd"]


def test_outage_stops_lookups_and_caches_nothing(tmp_path):
    s, fake = sizer(tmp_path, {"A": "DOWN", "B": org(3, None, "B")})
    assert s.assess("A") is None and s.assess("B") is None
    assert fake.calls == ["A"] and s.stopped and s.cache == {}            # both retried next run


def test_lookup_cap(tmp_path):
    s, fake = sizer(tmp_path, {"A": org(3, None, "A"), "B": org(3, None, "B")}, max_lookups=1)
    assert s.assess("A")["small"] and s.assess("B") is None and fake.calls == ["A"]


# ---- pipeline

def test_small_company_excluded_before_scoring_and_domain_filled(tmp_path):
    s, _ = sizer(tmp_path, {"Acme Analytics": org(4, "acme.io", "Acme Analytics"),
                            "Gamma AI": org(1200, "gamma.ai", "Gamma AI")})
    llm = FakeLLM(evaluations=[top_eval()])
    ws, stats = run(llm, tailor=False, company_sizer=s, open_checker=opened)
    assert stats["excluded_small_company"] == 1 and llm.calls["submit_evaluation"] == 1   # Acme never scored
    data = {r["Company"]: r for r in rows(ws).values()}
    acme, gamma = data["Acme Analytics"], data["Gamma AI"]
    assert acme["Processing Status"] == "excluded_small_company" and "Apollo: 4 employees" in acme["Error Details"]
    assert acme["Employees"] == "4" and acme["Match Score"] is None
    assert gamma["Employees"] == "1200" and gamma["Company Domain"] == "gamma.ai"
    shown = _master(load_workbook(ws.jobs_xlsx))
    assert [(m["Company"], m["Employees"], m["Company Domain"]) for m in shown] == [("Gamma AI", "1200", "gamma.ai")]


def test_namesake_never_excludes_or_fills_domain_in_the_pipeline(tmp_path):
    s, _ = sizer(tmp_path, {"Acme Analytics": org(4, "acme-namesake.io", "Acme Analytics", linkedin="other")})
    ws, stats = run(FakeLLM(evaluations=[top_eval(), top_eval()]), tailor=False, company_sizer=s, open_checker=opened)
    acme = next(r for r in rows(ws).values() if r["Company"] == "Acme Analytics")
    assert stats["excluded_small_company"] == 0 and acme["Employees"] == "unknown" and acme["Company Domain"] is None


def test_user_typed_domain_is_never_overwritten_and_drives_an_exact_lookup(tmp_path):
    ws, _ = run(FakeLLM(evaluations=[top_eval(), top_eval()]), tailor=False, open_checker=opened)  # no size check
    store = JobStore(ws.jobs_xlsx)
    for r in store.all():
        store.fill_company_field(r["Company"], "Company Domain", "mine.com")
    store.save()
    s, fake = sizer(tmp_path, {"mine.com": org(900, "other.com", "Mine")})
    run(FakeLLM(), tailor=False, company_sizer=s, skip_discovery=True, run_id="r2")
    assert fake.calls == ["mine.com", "mine.com"]                     # by domain, once per company
    data = [r for r in rows(ws).values() if r["Match Score"]]
    assert {r["Company Domain"] for r in data} == {"mine.com"} and {r["Employees"] for r in data} == {"900"}


def test_existing_master_rows_are_backfilled_and_small_ones_drop_out(tmp_path):
    ws, _ = run(FakeLLM(evaluations=[top_eval(), top_eval()]), tailor=False, open_checker=opened)  # before the rule
    assert len(_master(load_workbook(ws.jobs_xlsx))) == 2
    s, fake = sizer(tmp_path, {"Acme Analytics": org(2, "acme.io", "Acme Analytics")})   # Gamma: no record
    _, stats = run(FakeLLM(), tailor=False, company_sizer=s, skip_discovery=True, run_id="r2")
    assert stats["excluded_small_company"] == 1 and sorted(fake.calls) == ["Acme Analytics", "Gamma AI"]
    shown = _master(load_workbook(ws.jobs_xlsx))
    assert [(m["Company"], m["Employees"]) for m in shown] == [("Gamma AI", "unknown")]   # unknown is kept
    run(FakeLLM(), tailor=False, company_sizer=s, skip_discovery=True, run_id="r3")
    assert len(fake.calls) == 2                                            # nothing looked up again


def test_applied_and_not_working_rows_cost_no_lookups(tmp_path):
    ws, _ = run(FakeLLM(evaluations=[top_eval(), top_eval()]), tailor=False, open_checker=opened)
    store = JobStore(ws.jobs_xlsx)
    for r, choice in zip(store.all(), ("Applied", "Not Working")):
        store.ws.cell(row=store._rows[str(r["Job ID"])], column=store.col["Application Status"], value=choice)
    store.save()
    s, fake = sizer(tmp_path, {"Acme Analytics": org(2, "acme.io", "Acme Analytics")})
    _, stats = run(FakeLLM(), tailor=False, company_sizer=s, skip_discovery=True, run_id="r2")
    assert fake.calls == [] and stats["excluded_small_company"] == 0


def test_low_score_jobs_cost_one_lookup_each_before_scoring_only(tmp_path):
    s, fake = sizer(tmp_path, {})
    run(FakeLLM(evaluations=[evaluation((3, 3, 3, 3)), evaluation((3, 3, 3, 3))]), tailor=False, company_sizer=s)
    assert sorted(fake.calls) == ["Acme Analytics", "Gamma AI"]
