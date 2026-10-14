"""jobs.xlsx: the local job tracker and the cross-run dedup store (keyed by Job ID).

Rules:
- Existing rows are preserved; the workbook is loaded and updated in place, never recreated.
- Columns are located by header name, so reordered or extra user columns are tolerated.
- "Application Status" is the user's column: set to "Not Applied" on insert, never overwritten. So are
  "Company Domain", "Email Format" and "Contact Emails", which the user fills on the Master tab.
- Saves go to a temp file then atomically replace jobs.xlsx, keeping jobs.xlsx.bak.
"""

from __future__ import annotations

import os
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SHEET = "Jobs"
COLUMNS = [
    "Job ID", "Source Reference ID", "Job Title", "Company", "Location", "Job Type", "Work Mode", "Salary",
    "Job URL", "Job Description", "Required Skills", "Preferred Skills", "Matched Skills", "Missing Skills",
    "Match Score", "Application Status", "Date Discovered", "Resume Path", "Tailoring Report Path",
    "Processing Status", "Last Updated", "Error Details",
    "Company Domain", "Email Format", "Contact Emails",
]
COMPANY_FIELDS = ("Company Domain", "Email Format", "Contact Emails")   # one value per company
USER_OWNED = {"Application Status", *COMPANY_FIELDS}
DEFAULT_APPLICATION_STATUS = "Not Applied"
APPLICATION_CHOICES = ("Not Applied", "Applied", "Not Working")   # dropdown on Jobs + Master
MASTER_SHEET = "Master"
MASTER_EDITABLE = (("Applied?", "Application Status"), ("Company Domain", "Company Domain"),
                   ("Email Format", "Email Format"), ("Contact Emails", "Contact Emails"))   # (Master header, Jobs column)
LEGACY_VIEW_SHEETS = ("LinkedIn", "Indeed", "Naukri")   # replaced by Master; edits left on them are still synced once
LIST_SEP = "; "
MAX_CELL = 32000  # Excel's hard limit is 32,767 characters per cell
WIDTHS = {"Job ID": 18, "Job Title": 40, "Company": 30, "Job URL": 45, "Job Description": 60,
          "Resume Path": 60, "Tailoring Report Path": 60, "Error Details": 50}


class WorkbookLockedError(RuntimeError):
    """jobs.xlsx can't be written, usually because it's open in Excel."""


def _cell_value(value):
    if isinstance(value, (list, tuple)):
        value = LIST_SEP.join(str(v) for v in value if str(v).strip())
    if isinstance(value, str):
        value = value[:MAX_CELL]
    return value


def company_key(name) -> str:
    """Company name normalized for matching rows of the same company ('Acme  Analytics' == 'acme analytics')."""
    return " ".join(str(name or "").split()).casefold()


def clean_domain(value) -> str | None:
    """'https://www.Fractal.ai/careers' -> 'fractal.ai'."""
    v = str(value or "").strip().lower()
    v = re.sub(r"^[a-z]+://", "", v).split("/")[0].split("?")[0]
    v = v.removeprefix("www.")
    return v or None


def split_list(value) -> list[str]:
    if not value:
        return []
    return [v.strip() for v in str(value).split(LIST_SEP.strip()) if v.strip()]


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


class JobStore:
    def __init__(self, path: Path):
        self.path = Path(path)
        if self.path.exists():
            try:
                self.wb = load_workbook(self.path)
            except PermissionError as exc:
                raise WorkbookLockedError(f"cannot open {self.path}; close it in Excel and re-run") from exc
            self.ws = self.wb[SHEET] if SHEET in self.wb.sheetnames else self.wb.create_sheet(SHEET)
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = SHEET
        self._ensure_headers()
        self._rows = self._index_rows()
        self.view_edits_synced = self._sync_view_edits()

    # ---------------------------------------------------------------- structure

    def _ensure_headers(self) -> None:
        ws = self.ws
        existing = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        if not any(existing):
            existing = []
        for name in COLUMNS:
            if name not in existing:
                existing.append(name)
                cell = ws.cell(row=1, column=len(existing), value=name)
                cell.font = Font(bold=True)
                letter = cell.column_letter
                ws.column_dimensions[letter].width = WIDTHS.get(name, 18)
        ws.freeze_panes = "B2"
        self.col = {c.value: c.column for c in ws[1] if c.value}

    def _index_rows(self) -> dict[str, int]:
        rows, jcol = {}, self.col["Job ID"]
        for r in range(2, self.ws.max_row + 1):
            jid = self.ws.cell(row=r, column=jcol).value
            if jid:
                rows[str(jid)] = r
        return rows

    # ---------------------------------------------------------------- read

    def __contains__(self, job_id: str) -> bool:
        return job_id in self._rows

    def __len__(self) -> int:
        return len(self._rows)

    def get(self, job_id: str) -> dict | None:
        r = self._rows.get(job_id)
        if r is None:
            return None
        return {name: self.ws.cell(row=r, column=c).value for name, c in self.col.items()}

    def all(self) -> list[dict]:
        return [self.get(j) for j in self._rows]

    def dedup_index(self) -> dict:
        """Index in the shape scrape_jobs.deduplicate() expects."""
        index = {}
        for row in self.all():
            index[str(row["Job ID"])] = {
                "ref_id": row["Source Reference ID"], "title": row["Job Title"], "company": row["Company"],
                "link": row["Job URL"], "first_seen_at": row["Date Discovered"], "last_seen_at": row["Last Updated"],
                "job_status": row["Processing Status"],
            }
        return index

    # ---------------------------------------------------------------- write

    def upsert(self, job_id: str, values: dict) -> str:
        """Insert or update one row by Job ID. Returns 'inserted' or 'updated'."""
        unknown = set(values) - set(self.col)
        if unknown:
            raise KeyError(f"unknown columns: {sorted(unknown)}")
        r = self._rows.get(job_id)
        action = "updated"
        if r is None:
            r = self.ws.max_row + 1
            self._rows[job_id] = r
            action = "inserted"
            self.ws.cell(row=r, column=self.col["Job ID"], value=job_id)
            self.ws.cell(row=r, column=self.col["Application Status"], value=DEFAULT_APPLICATION_STATUS)
        for name, value in values.items():
            if name == "Job ID" or (name in USER_OWNED and action == "updated"):
                continue
            cell = self.ws.cell(row=r, column=self.col[name])
            cell.value = _cell_value(value)  # ws.cell(value=None) would silently keep the old value
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # scraped text must never become a formula
            if name in ("Job Description", "Error Details"):
                cell.alignment = Alignment(wrap_text=False, vertical="top")
        self.ws.cell(row=r, column=self.col["Last Updated"], value=now_iso())
        return action

    def _sync_view_edits(self) -> int:
        """Copy what the user typed on the Master tab back into Jobs.

        Master is rebuilt on every save, so without this an edit made there would be lost. Only cells
        the user changed since the tab was built (vs. its hidden '_built:<column>' copy) are copied, so
        a stale tab never overwrites a newer edit made on the Jobs sheet. Applied? goes to that job's
        row; the company fields go to every row of the same company. The old LinkedIn/Indeed/Naukri
        tabs (Applied? + '_built') are read too, so choices left on them survive the switch to Master.
        """
        synced = 0
        for name in (MASTER_SHEET, *LEGACY_VIEW_SHEETS):
            if name not in self.wb.sheetnames:
                continue
            vs = self.wb[name]
            head = {c.value: c.column for c in vs[4] if c.value}
            if "Job ID" not in head:
                continue
            for view_col, field in MASTER_EDITABLE:
                built_col = head.get(f"_built:{view_col}") or (head.get("_built") if view_col == "Applied?" else None)
                if view_col not in head or not built_col:
                    continue
                for r in range(5, vs.max_row + 1):
                    jid = vs.cell(row=r, column=head["Job ID"]).value
                    val = vs.cell(row=r, column=head[view_col]).value
                    built = vs.cell(row=r, column=built_col).value
                    if not jid or str(jid) not in self._rows or val == built:
                        continue
                    if field == "Application Status":
                        if val not in APPLICATION_CHOICES:
                            continue
                        targets = [self._rows[str(jid)]]
                    else:
                        val = clean_domain(val) if field == "Company Domain" else (str(val or "").strip() or None)
                        key = company_key(self.ws.cell(row=self._rows[str(jid)], column=self.col["Company"]).value)
                        targets = [row for row in self._rows.values()
                                   if company_key(self.ws.cell(row=row, column=self.col["Company"]).value) == key]
                    changed = False
                    for row in targets:
                        cell = self.ws.cell(row=row, column=self.col[field])
                        if cell.value != val:
                            cell.value = val
                            changed = True
                    synced += changed
        return synced

    def _apply_dropdown(self) -> None:
        letter = get_column_letter(self.col["Application Status"])
        add_choice_dropdown(self.ws, f"{letter}2:{letter}{max(2, self.ws.max_row) + 5000}")

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.ws.auto_filter.ref = self.ws.dimensions  # filter dropdowns on every column
        self._apply_dropdown()
        from dashboard import build_dashboard     # summary sheet, rebuilt from the rows on every save
        from views import build_views              # Master: every open 8+ opening, all platforms
        rows = self.all()
        build_dashboard(self.wb, rows)
        build_views(self.wb, rows, self.path.parent / "logs")
        self.wb.active = self.wb.sheetnames.index(SHEET)
        tmp = self.path.with_name(self.path.stem + ".tmp.xlsx")
        try:
            self.wb.save(tmp)
            if self.path.exists():
                shutil.copy2(self.path, self.path.with_suffix(".xlsx.bak"))
            os.replace(tmp, self.path)
        except PermissionError as exc:
            tmp.unlink(missing_ok=True)
            raise WorkbookLockedError(f"cannot write {self.path}; close it in Excel and re-run") from exc


def add_choice_dropdown(ws, cell_range: str) -> None:
    """Applied / Not Applied / Not Working dropdown on a range (replaces any earlier one on that column)."""
    col = "".join(ch for ch in cell_range.split(":")[0] if ch.isalpha())
    ws.data_validations.dataValidation = [
        dv for dv in ws.data_validations.dataValidation
        if not any("".join(ch for ch in str(rng).split(":")[0] if ch.isalpha()) == col for rng in dv.sqref.ranges)]
    dv = DataValidation(type="list", formula1='"' + ",".join(APPLICATION_CHOICES) + '"', allow_blank=True,
                        showDropDown=False, showErrorMessage=True, errorTitle="Application status",
                        error="Choose Applied, Not Applied or Not Working.")
    dv.add(cell_range)
    ws.add_data_validation(dv)
